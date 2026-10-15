#!/usr/bin/env python3
"""Render docs/BUILD_AND_VALIDATION_GUIDE.md into the Word copy of the guide.

The markdown file is the single source of truth; this script only formats it. Run:

    python -m pip install python-docx
    python docs/tools/build_guide_docx.py

Requires Python 3.9+ and `python-docx`. It writes
docs/JARVIS_Build_Test_Evaluation_Guide.docx next to the markdown.
"""

from __future__ import annotations

import pathlib
import re
import sys

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
except ImportError:  # pragma: no cover - guidance for a missing dev dependency
    sys.exit("python-docx is required: python -m pip install python-docx")

DOCS = pathlib.Path(__file__).resolve().parent.parent
SOURCE = DOCS / "BUILD_AND_VALIDATION_GUIDE.md"
TARGET = DOCS / "JARVIS_Build_Test_Evaluation_Guide.docx"

LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
BOLD = re.compile(r"\*\*([^*]+)\*\*")
CODE = re.compile(r"`([^`]+)`")


def plain(text: str) -> str:
    """Markdown -> readable plain text (Word copy keeps the wording, not the markup)."""
    text = LINK.sub(lambda m: f"{m.group(1)} ({m.group(2)})", text)
    text = BOLD.sub(r"\1", text)
    text = CODE.sub(r"\1", text)
    return text.strip()


def add_runs(paragraph, text: str) -> None:
    """Add text to a paragraph, keeping bold spans and rendering code in monospace."""
    for bold_match in BOLD.finditer(text):
        pass  # handled below by a simple token walk

    token = re.compile(r"\*\*([^*]+)\*\*|`([^`]+)`")
    position = 0
    for match in token.finditer(text):
        if match.start() > position:
            paragraph.add_run(plain(text[position:match.start()]))
        if match.group(1) is not None:
            paragraph.add_run(match.group(1)).bold = True
        else:
            run = paragraph.add_run(match.group(2))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        position = match.end()
    if position < len(text):
        paragraph.add_run(plain(text[position:]))


def flush_table(document, rows: list[list[str]]) -> None:
    if not rows:
        return
    header, *body = rows
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for index, cell in enumerate(header):
        paragraph = table.rows[0].cells[index].paragraphs[0]
        run = paragraph.add_run(plain(cell))
        run.bold = True
    for row in body:
        cells = table.add_row().cells
        for index, value in enumerate(row[: len(header)]):
            add_runs(cells[index].paragraphs[0], value)
    document.add_paragraph()


def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def build() -> pathlib.Path:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    document.core_properties.title = "JARVIS - Build, Test, and Evaluation Guide"
    document.core_properties.comments = (
        "Generated from docs/BUILD_AND_VALIDATION_GUIDE.md by "
        "docs/tools/build_guide_docx.py - edit the markdown, not this file."
    )

    table_rows: list[list[str]] = []
    code_block: list[str] = []
    in_code = False

    for line in lines:
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Pt(18)
                for code_line in code_block:
                    run = paragraph.add_run(code_line + "\n")
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
                code_block = []
            in_code = not in_code
            continue
        if in_code:
            code_block.append(line)
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            if set(stripped) <= set("|-: "):
                continue  # separator row
            table_rows.append(split_row(stripped))
            continue
        if table_rows:
            flush_table(document, table_rows)
            table_rows = []

        if not stripped:
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            document.add_heading(plain(stripped.lstrip("# ")), level=min(level, 4))
        elif stripped.startswith(("- ", "* ")):
            paragraph = document.add_paragraph(style="List Bullet")
            add_runs(paragraph, stripped[2:])
        elif re.match(r"^\d+\.\s", stripped):
            paragraph = document.add_paragraph(style="List Number")
            add_runs(paragraph, re.sub(r"^\d+\.\s", "", stripped))
        elif stripped.startswith(">"):
            paragraph = document.add_paragraph(style="Intense Quote")
            add_runs(paragraph, stripped.lstrip("> "))
        else:
            paragraph = document.add_paragraph()
            add_runs(paragraph, stripped)

    if table_rows:
        flush_table(document, table_rows)
    if code_block:
        paragraph = document.add_paragraph()
        for code_line in code_block:
            run = paragraph.add_run(code_line + "\n")
            run.font.name = "Consolas"
            run.font.size = Pt(9)

    footer = document.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = footer.add_run(
        f"Generated from {SOURCE.name} by docs/tools/build_guide_docx.py. "
        "The markdown file is the authoritative version."
    )
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    document.save(TARGET)
    return TARGET


if __name__ == "__main__":
    written = build()
    print(f"wrote {written.relative_to(DOCS.parent)} ({written.stat().st_size} bytes)")
